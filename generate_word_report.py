"""
Academic Word Document Generator for CollegeCare AI
Generates an official formatted .docx capstone report.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report():
    doc = Document()

    # 1. Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Base Style Formatting
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Read Markdown source
    md_path = "FINAL_PROJECT_REPORT.md"
    if not os.path.exists(md_path):
        print("FINAL_PROJECT_REPORT.md not found!")
        return

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Parse rows and cols
        headers = [c.strip() for c in table_rows[0].strip('| \n').split('|')]
        data = []
        for r in table_rows[1:]:
            # Skip separator row like |---|---|
            if set(r.replace('|', '').replace('-', '').replace(':', '').strip()) == set():
                continue
            cols = [c.strip() for c in r.strip('| \n').split('|')]
            data.append(cols)

        if not headers or not data:
            in_table = False
            table_rows = []
            return

        col_count = len(headers)
        table = doc.add_table(rows=len(data) + 1, cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            if i < len(hdr_cells):
                hdr_cells[i].text = h
                set_cell_background(hdr_cells[i], "1E293B")
                p = hdr_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.bold = True
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Data Rows
        for row_idx, row_data in enumerate(data):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                if col_idx < len(row_cells):
                    row_cells[col_idx].text = cell_value
                    set_cell_background(row_cells[col_idx], bg_color)
                    p = row_cells[col_idx].paragraphs[0]
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(10)

        doc.add_paragraph()  # spacing
        in_table = False
        table_rows = []

    for line in lines:
        stripped = line.strip()

        # Handle Table
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_rows.append(stripped)
            continue
        elif in_table:
            flush_table()

        # Handle Page Break
        if "page-break-after: always" in line or stripped == "---":
            if "page-break-after: always" in line:
                doc.add_page_break()
                continue
            elif stripped == "---":
                continue

        # Skip raw HTML tags
        if stripped.startswith("<") and stripped.endswith(">"):
            continue

        # Skip empty lines
        if not stripped:
            continue

        # Heading 1 (Chapters or Major Sections)
        if stripped.startswith("# "):
            text = stripped[2:].strip().replace("**", "")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if "A PROJECT REPORT" in text or "COLLEGECARE" in text else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        # Heading 2 (Subsections or Chapter titles)
        elif stripped.startswith("## "):
            text = stripped[3:].strip().replace("**", "")
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # Heading 3 (Detailed topics)
        elif stripped.startswith("### "):
            text = stripped[4:].strip().replace("**", "")
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        # Bullet List
        elif stripped.startswith("* ") or stripped.startswith("- "):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            # Simple bold handler inside bullets
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = 'Times New Roman'

        # Numbered List
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = 'Times New Roman'

        # Standard Paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = 'Times New Roman'

    if in_table:
        flush_table()

    # Save to workspace
    out_workspace = "CollegeCare_AI_Final_Project_Report.docx"
    doc.save(out_workspace)
    print(f"Saved: {out_workspace}")

    # Also save directly to Desktop for one-click access
    desktop_dir = r"C:\Users\Shravani .D. Thukrul\OneDrive\Desktop"
    if os.path.exists(desktop_dir):
        out_desktop = os.path.join(desktop_dir, "CollegeCare_AI_Final_Project_Report.docx")
        doc.save(out_desktop)
        print(f"Saved to Desktop: {out_desktop}")

if __name__ == "__main__":
    create_report()
